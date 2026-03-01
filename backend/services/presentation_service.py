"""
Presentation Upload & Slide Conversion Service.
Converts PDF/PPTX to individual slide images for webinar rendering.
"""
import os
import io
import uuid
import logging
import tempfile
from typing import List

from services.object_storage import put_object, get_object

logger = logging.getLogger(__name__)


def convert_pdf_to_images(pdf_data: bytes) -> List[bytes]:
    """Convert PDF bytes to list of PNG image bytes (one per page)."""
    from pdf2image import convert_from_bytes
    images = convert_from_bytes(pdf_data, dpi=150, fmt='png')
    result = []
    for img in images:
        buf = io.BytesIO()
        img.save(buf, format='PNG', optimize=True)
        result.append(buf.getvalue())
    return result


def convert_pptx_to_images(pptx_data: bytes) -> List[bytes]:
    """Convert PPTX to slide images. Uses python-pptx to extract slide content as simplified images."""
    from pptx import Presentation
    from PIL import Image, ImageDraw, ImageFont

    prs = Presentation(io.BytesIO(pptx_data))
    slide_width = prs.slide_width.inches
    slide_height = prs.slide_height.inches
    scale = 150
    w = int(slide_width * scale)
    h = int(slide_height * scale)

    result = []
    for slide in prs.slides:
        img = Image.new('RGB', (w, h), color=(20, 20, 35))
        draw = ImageDraw.Draw(img)

        for shape in slide.shapes:
            if shape.has_text_frame:
                x = int(shape.left / 914400 * scale) if shape.left else 40
                y = int(shape.top / 914400 * scale) if shape.top else 40
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        font_size = 16
                        if para.runs:
                            run = para.runs[0]
                            if run.font.size:
                                font_size = max(10, int(run.font.size.pt * 1.2))
                        try:
                            font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", min(font_size, 48))
                        except:
                            font = ImageFont.load_default()
                        draw.text((x, y), text, fill='white', font=font)
                        y += font_size + 8

        buf = io.BytesIO()
        img.save(buf, format='PNG', optimize=True)
        result.append(buf.getvalue())
    return result


async def process_presentation(file_data: bytes, filename: str, webinar_id: str, user_id: str) -> dict:
    """Process uploaded presentation and store slide images in cloud."""
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''

    if ext == 'pdf':
        slides = convert_pdf_to_images(file_data)
    elif ext in ('pptx', 'ppt'):
        slides = convert_pptx_to_images(file_data)
    else:
        raise ValueError(f"Unsupported format: {ext}. Use PDF or PPTX.")

    slide_paths = []
    for i, slide_data in enumerate(slides):
        path = f"ai-karau/presentations/{webinar_id}/{uuid.uuid4()}_slide_{i}.png"
        result = put_object(path, slide_data, "image/png")
        slide_paths.append(result["path"])

    return {
        "webinar_id": webinar_id,
        "filename": filename,
        "total_slides": len(slides),
        "slide_paths": slide_paths,
        "uploaded_by": user_id
    }
